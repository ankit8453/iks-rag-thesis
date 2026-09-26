"""Convert one of our Markdown deliverables into a readable .docx.

Written for RESEARCH_WRITEUP_*.md but kept general: it handles the subset of
Markdown we actually use — ATX headings, paragraphs, bullet and numbered lists,
pipe tables, fenced code blocks (kept monospace for the ASCII architecture
diagram), block quotes, horizontal rules, and inline **bold** / *italic* / `code`.

Usage:  python scripts/md_to_docx.py RESEARCH_WRITEUP_2026-09-20.md
        python scripts/md_to_docx.py in.md out.docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x66, 0x66, 0x66)

HEADING_SIZES = {1: 16, 2: 13, 3: 11.5}

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\n]+\*(?![\w*]))", re.DOTALL)


def style_doc(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(54)
        s.left_margin = s.right_margin = Pt(54)


def add_inline(p, text: str, *, size=10.5, bold=False, italic=False, font=BODY_FONT):
    """Add text to a paragraph, honouring **bold**, *italic* and `code`."""
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i, f = bold, italic, font
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, f = part[1:-1], MONO_FONT
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        r = p.add_run(part)
        r.font.name = f
        r.font.size = Pt(size if f != MONO_FONT else size - 1)
        r.bold = b
        r.italic = i
    return p


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s:|-]+\|[\s:|-]*", line.strip())) and "-" in line


def add_table(doc, rows: list[list[str]]) -> None:
    t = doc.add_table(rows=0, cols=max(len(r) for r in rows))
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            if ci >= len(cells):
                break
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, val, size=9, bold=(ri == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_doc(doc)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code block -> monospace, preserved verbatim
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for code_line in block:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(code_line or " ")
                r.font.name = MONO_FONT
                r.font.size = Pt(7.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # pipe table
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = doc.add_paragraph()
            r = p.add_run("_" * 96)
            r.font.size = Pt(7)
            r.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        # heading
        m = re.match(r"(#{1,6})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, m.group(2), size=HEADING_SIZES.get(level, 10.5), bold=True)
            for r in p.runs:
                r.font.color.rgb = ACCENT
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # block quote
        if stripped.startswith(">"):
            body = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                body.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, " ".join(body), italic=True)
            for r in p.runs:
                r.font.color.rgb = GREY
            continue

        # bullet / numbered list item
        m = re.match(r"[-*]\s+(.*)", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(1))
            i += 1
            continue
        m = re.match(r"\d+\.\s+(.*)", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(1))
            i += 1
            continue

        # paragraph: join until a blank line or a structural marker
        body = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```", "- ", "* "))
                    or re.match(r"\d+\.\s", nxt)
                    or re.fullmatch(r"-{3,}|\*{3,}|_{3,}", nxt)):
                break
            body.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(body))

    doc.save(out_path)
    print(f"wrote {out_path}")


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
