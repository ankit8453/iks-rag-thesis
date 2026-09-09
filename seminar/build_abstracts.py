"""Build the ICSSR seminar ABSTRACT drafts (<=250 words, TNR 14pt .docx) + a markdown copy.

Deliverable for 8 Sep 2026: title + 250-word abstract (+ registration form).

Each abstract follows the standard research-abstract structure:
  background -> problem/gap -> objective -> methodology -> (preliminary) results -> conclusion.
No system walkthrough, no interim metrics, no implementation details.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent

SEMINAR = ("International Seminar on Revitalising Indigenous Knowledge Systems: "
           "Tribal Heritage, Cultural Continuity and Social Transformation in India")
VENUE = ("Centre for Tribal Studies, Iswar Saran Degree College (University of Allahabad), "
         "Prayagraj  |  11-12 September 2026  |  Sponsored by ICSSR")
SUBTHEME = ("Sub-theme: Digital Preservation, Artificial Intelligence, and "
            "Documentation of Indigenous Knowledge")
AUTHORS = [
    "Ankit Pawar",
    "M.Tech (Computer Science & Engineering), PDPM Indian Institute of Information "
    "Technology, Design and Manufacturing, Jabalpur",
    "Dr. Akshay Pandey",
    "Department of Computer Science & Engineering, PDPM IIITDM Jabalpur",
]
KEYWORDS = ("Indian Knowledge Systems; Vrikshayurveda; digital preservation; artificial "
            "intelligence; retrieval-augmented generation; agricultural advisory")

SIX = ("Vrikshayurveda, Brihat Samhita, Krishi Parashara, Upavanavinoda, Kashyapiya "
       "Krishisukti and Vishvavallabha")

# ---------------------------------------------------------------------------
# DRAFT A - primary (balanced: preservation + application + faithfulness)
# ---------------------------------------------------------------------------
DA = (
    f"India's classical agricultural treatises - {SIX} - embody a sustainable, organic body "
    "of knowledge on plant health, soil and cultivation, yet remain largely unused in "
    "current practice. Three barriers persist: the texts lack a structured, "
    "searchable digital form; they describe plant disorders by observed symptoms rather than "
    "the disease taxonomy of modern diagnosis; and current AI-based advisory tools neither "
    "draw on them nor guarantee faithful, source-attributed recommendations. This study "
    "proposes an Indian Knowledge Systems (IKS)-grounded multimodal advisory system that "
    "connects image-based diagnosis of plant disease and soil condition with treatment "
    "knowledge from these texts. The methodology comprises a verse-level, "
    "metadata-annotated digital corpus of the six treatises; vision models for disease and "
    "soil-property recognition, with visual explanation used to verify that predictions rest "
    "on the lesion rather than the background; a symptom-based semantic bridge that "
    "re-expresses modern diagnostic labels in the descriptive vocabulary of the texts, "
    "enabling cautious advice even for crops outside the training data; and "
    "retrieval-augmented generation constrained to retrieved passages, with verse-level "
    "citation and abstention where the texts are silent. Evaluation combines retrieval "
    "metrics, faithfulness measures and domain-expert assessment, including quantification "
    "of hallucination against IKS sources. Preliminary results indicate that the "
    "symptom-based bridge substantially improves retrieval relevance over direct label "
    "matching, and that the grounded generator produces source-cited recommendations "
    "without fabrication. The work contributes the first searchable digital corpus of these "
    "treatises and a reproducible framework for the faithful digital preservation and field "
    "application of indigenous agricultural knowledge."
)

# ---------------------------------------------------------------------------
# DRAFT B - digital preservation & documentation emphasis (sub-theme wording)
# ---------------------------------------------------------------------------
DB = (
    f"Indigenous agricultural knowledge in India's classical treatises - {SIX} - faces "
    "erosion: it survives only in scattered printed and scanned editions, is organised by "
    "observed symptom rather than modern disease categories, and is therefore invisible to "
    "the digital tools on which farmers and extension workers now rely. This study addresses "
    "the gap between documentation and use by proposing an Indian Knowledge Systems "
    "(IKS)-grounded multimodal advisory system. Its objectives are to build the first "
    "structured, verse-level annotated digital corpus of the six treatises; to connect "
    "image-based recognition of plant disease and soil condition with the treatments the "
    "texts prescribe; and to ensure that every recommendation remains traceable to its "
    "source. The approach combines vision models whose attention is verified through visual "
    "explanation, a symptom-based semantic bridge that renders modern diagnostic labels in "
    "the texts' own descriptive vocabulary, and retrieval-augmented generation restricted to "
    "retrieved passages with verse-level citation and abstention where the texts are silent; "
    "uncertain cases are routed to expert review. Evaluation integrates retrieval metrics, "
    "faithfulness measures and domain-expert judgement. Preliminary findings show that "
    "symptom-based bridging markedly improves retrieval of relevant passages compared with "
    "direct label matching, and that grounded generation yields cited, non-fabricated "
    "recommendations. The study contributes a reproducible framework in which digital "
    "preservation and practical application reinforce one another, offering a model for "
    "documenting Indigenous Knowledge Systems as living, verifiable knowledge rather than "
    "static archives."
)

# ---------------------------------------------------------------------------
# DRAFT C - faithful / trustworthy AI emphasis
# ---------------------------------------------------------------------------
DC = (
    "Generative AI is increasingly used for agricultural advice, but large language models "
    "routinely fabricate facts and attribute them to non-existent sources - a critical risk "
    "when the knowledge in question is indigenous, textual and difficult for users to "
    f"verify. India's classical agricultural treatises - {SIX} - contain sustainable "
    "plant-care and soil knowledge organised by observed symptom rather than modern disease "
    "names, and have no structured digital presence. This study proposes an Indian Knowledge "
    "Systems (IKS)-grounded multimodal advisory system in which every recommendation is "
    "constrained to, and cited from, the texts themselves. Its objectives are to construct a "
    "verse-level annotated corpus of the six treatises; to connect image-based diagnosis of "
    "plant disease and soil condition with classical prescriptions through a symptom-based "
    "semantic bridge; and to measure, for the first time, hallucination of AI against IKS "
    "sources. Methods include vision models whose attention is verified through visual "
    "explanation, hybrid retrieval, citation verification, abstention where the texts are "
    "silent, and calibrated confidence for crops outside the training data. Evaluation "
    "combines retrieval metrics, automated faithfulness scoring and domain-expert "
    "assessment. Preliminary results indicate substantially improved retrieval relevance "
    "from symptom-based bridging and source-cited generation free of fabricated citations. "
    "The work offers a verifiable framework for applying AI to Indigenous Knowledge Systems "
    "without distorting them, with implications for digital preservation, extension "
    "services and the responsible use of AI in cultural-heritage domains."
)

DRAFTS = [
    ("DRAFT A (RECOMMENDED) - balanced: preservation, application, faithfulness",
     "An IKS-Grounded Multimodal Advisory System: Connecting Image-Based Plant Diagnosis "
     "with Classical Indian Agricultural Texts", DA),
    ("DRAFT B - digital preservation & documentation emphasis",
     "Digital Preservation and Application of Indian Agricultural Knowledge: A Multimodal, "
     "Retrieval-Grounded Advisory System over Classical Texts", DB),
    ("DRAFT C - faithful / trustworthy AI emphasis",
     "Faithful Artificial Intelligence for Indigenous Knowledge: Grounding Agricultural "
     "Advice in Classical Indian Texts", DC),
]


def wc(text: str) -> int:
    return len(re.findall(r"\S+", text))


def style_doc(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), "Times New Roman")


def para(doc, text, *, bold=False, italic=False, align=None, size=14, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    return p


def header_block(doc):
    para(doc, SEMINAR, bold=True, align="center")
    para(doc, VENUE, italic=True, align="center", size=12)
    para(doc, SUBTHEME, bold=True, align="center", size=13, space_after=14)


def abstract_block(doc, title, body, label=None, show_count=False):
    if label:
        para(doc, label, bold=True, italic=True, size=12, space_after=4)
    para(doc, title, bold=True, align="center", space_after=8)
    for i, a in enumerate(AUTHORS):
        para(doc, a, bold=(i % 2 == 0), align="center", size=12, space_after=2)
    para(doc, "", space_after=4)
    para(doc, "Abstract", bold=True, space_after=4)
    para(doc, body, align="justify", space_after=8)
    para(doc, "Keywords: " + KEYWORDS, italic=True, size=13, space_after=4)
    if show_count:
        para(doc, f"[word count: {wc(body)} / limit 250]", italic=True, size=11, space_after=14)


def save_doc(doc, path: Path) -> Path:
    """Save; if Word has the file open (PermissionError), save as *_vN instead."""
    candidates = [path] + [path.with_name(f"{path.stem}_v{n}{path.suffix}") for n in range(2, 10)]
    for cand in candidates:
        try:
            doc.save(cand)
            if cand is not path:
                print(f"  ! {path.name} is open in Word - saved as {cand.name} instead")
            return cand
        except PermissionError:
            continue
    raise PermissionError(f"could not save {path.name}: close it in Word and re-run")


def main() -> None:
    doc = Document()
    style_doc(doc)
    header_block(doc)
    para(doc, "Three candidate abstracts - same thesis, three emphases. Choose one.",
         italic=True, size=12, space_after=14)
    for i, (label, title, body) in enumerate(DRAFTS):
        abstract_block(doc, title, body, label, show_count=True)
        if i < len(DRAFTS) - 1:
            doc.add_page_break()
    save_doc(doc, HERE / "ICSSR_Abstract_DRAFTS_2026-09-08.docx")

    for letter, (label, title, body) in zip("ABC", DRAFTS):
        d = Document()
        style_doc(d)
        header_block(d)
        abstract_block(d, title, body)
        save_doc(d, HERE / f"ICSSR_Abstract_SUBMIT_Draft{letter}.docx")

    md = [f"# ICSSR Seminar - Abstract Drafts (2026-09-08)\n\n**{SEMINAR}**  \n*{VENUE}*  \n"
          f"**{SUBTHEME}**\n"]
    for label, title, body in DRAFTS:
        md.append(f"\n---\n\n## {label}\n\n### {title}\n\n"
                  + "  \n".join(AUTHORS)
                  + f"\n\n**Abstract.** {body}\n\n*Keywords:* {KEYWORDS}\n\n"
                  f"`word count: {wc(body)} / 250`\n")
    (HERE / "ICSSR_Abstract_DRAFTS_2026-09-08.md").write_text("".join(md), encoding="utf-8")

    for label, title, body in DRAFTS:
        flag = "" if wc(body) <= 250 else "   <-- OVER 250, trim"
        print(f"{wc(body):4} words | {label}{flag}")


if __name__ == "__main__":
    main()
