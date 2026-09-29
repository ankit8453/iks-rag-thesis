"""Build the chapter .docx with the two block diagrams as native Word tables.

Why tables and not pictures. The chapter is submitted as a Word file, and a reviewer can
tell a pasted image from an object built in Word. Drawn shapes with connector arrows
cannot be generated reliably from code, but a block diagram laid out as a table can:
shaded, bordered cells for the boxes, arrows as text in border-less cells between them.
The result is editable in Word and looks like it was made there — because it was.

How it works. The Markdown is converted with scripts/md_to_docx.py (which inserts the
PNG figures). Then each picture paragraph is replaced, in place, by the corresponding
table diagram; the italic caption paragraph that follows is kept.

Usage:  python paper/icssr_chapter/build_chapter_docx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from scripts.md_to_docx import convert  # noqa: E402

MD = HERE / "CHAPTER_DRAFT_v1.md"
# Optional output name on the command line, so a build can proceed while the canonical
# file is open in Word (Word locks it; python-docx then cannot overwrite it).
OUT = HERE / (sys.argv[1] if len(sys.argv) > 1 else "CHAPTER_DRAFT_v1.docx")

FONT = "Times New Roman"
INK = RGBColor(0x1F, 0x2A, 0x37)
BLUE = "DCE6F1"      # model / modern boxes
CREAM = "F4EFE6"     # input / classical boxes
GREEN = "E2F0D9"     # grounded answer / overlap
RED = "F8E1E1"       # refusal
LINE = "1F4E79"


# --------------------------------------------------------------------------- cell styling
def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(cell, *, on: bool, color: str = LINE, size: int = 8) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if on:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color); el.set(qn("w:space"), "0")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def _vcenter(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign"); v.set(qn("w:val"), "center"); tcPr.append(v)


def _write(cell, title: str = "", body: str = "", *, tsize=9.5, bsize=8, color=INK,
           align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    if title:
        r = p.add_run(title); r.bold = True; r.font.size = Pt(tsize); r.font.name = FONT
        r.font.color.rgb = color
    if body:
        if title:
            p = cell.add_paragraph(); p.alignment = align
            p.paragraph_format.space_after = Pt(1)
        r = p.add_run(body); r.font.size = Pt(bsize); r.font.name = FONT; r.font.color.rgb = color


def _box(cell, title, body="", fill=BLUE, **kw) -> None:
    _shade(cell, fill); _borders(cell, on=True); _vcenter(cell); _write(cell, title, body, **kw)


def _arrow(cell, text="→", note="", *, color=INK, size=14) -> None:
    _borders(cell, on=False); _vcenter(cell)
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text); r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = color
    if note:
        p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(note); r2.font.size = Pt(7); r2.font.name = FONT; r2.italic = True
        r2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def _blank(cell) -> None:
    _borders(cell, on=False); cell.text = ""


def _set_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def _new_table(doc, rows: int, cols: int):
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # no table-level borders; each cell decides
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblPr.append(b)
    return t


# --------------------------------------------------------------------------- Figure 1
def figure1(doc):
    """Inputs -> models -> bridge -> corpus -> outcome, as a 7-column grid.

        A input | a arrow | B model | b arrow | C bridge / corpus | c arrow | D outcome
    """
    t = _new_table(doc, rows=5, cols=7)
    c = lambda r, k: t.cell(r, k)  # noqa: E731
    green = RGBColor(0x2F, 0x6F, 0x4E); red = RGBColor(0x9B, 0x3B, 0x3B)

    # rows 0-1: the two photographs -> the two models -> the bridge (spanning both rows)
    _box(c(0, 0), "Leaf photograph", "from a phone camera", CREAM)
    _arrow(c(0, 1))
    _box(c(0, 2), "Disease model", "EfficientNet-B4, retrained on leaf crops; heat-map checked")
    _arrow(c(0, 3), note="disease label")
    _box(c(1, 0), "Soil photograph", "of the field surface", CREAM)
    _arrow(c(1, 1))
    _box(c(1, 2), "Soil model", "EfficientNet-B0: soil type, moisture, texture")
    _arrow(c(1, 3), note="soil reading")
    bridge = c(0, 4).merge(c(1, 4))
    _box(bridge, "Symptom bridge",
         "Llama 3.1 rewrites the label:\n“Apple scab” becomes “dark rough corky patches spreading over the leaves”")
    _blank(c(0, 5).merge(c(1, 5))); _blank(c(0, 6).merge(c(1, 6)))

    # row 2: crop name feeds the bridge as context; the query goes down to the corpus
    _box(c(2, 0), "Crop name", "typed by the farmer", CREAM)
    _arrow(c(2, 1).merge(c(2, 3)), "→", note="crop, as background context only")
    _arrow(c(2, 4), "↓", note="query, in the texts’ own words")
    _blank(c(2, 5)); _blank(c(2, 6))

    # rows 3-4: the corpus (spanning both) feeds either outcome
    for k in range(4):
        _blank(c(3, k)); _blank(c(4, k))
    corpus = c(3, 4).merge(c(4, 4))
    _box(corpus, "Digital corpus",
         "270 passages, 7 sources: six classical treatises + one modern manual; search and re-ranking")
    _arrow(c(3, 5), note="passages found", color=green)
    _box(c(3, 6), "Grounded answer", "cites [book, chapter, verse] for every claim", GREEN)
    _arrow(c(4, 5), note="nothing fits", color=red)
    _box(c(4, 6), "Honest refusal", "“the texts do not contain enough information”", RED)

    _set_widths(t, [2.6, 1.1, 3.3, 1.3, 3.6, 1.5, 3.2])
    return t


# --------------------------------------------------------------------------- Figure 2
def figure2(doc):
    """Classical (cause) | overlap | modern (appearance), with the finding underneath."""
    t = _new_table(doc, rows=4, cols=3)
    c = lambda r, k: t.cell(r, k)  # noqa: E731
    brown = "8A6D3B"

    _box(c(0, 0), "Classical treatises organise plant disorder by CAUSE",
         "wind · bile · phlegm imbalance · over-watering · over-manuring · insects · wounds · fire · unhealthy soil",
         CREAM); _borders(c(0, 0), on=True, color=brown)
    _box(c(0, 2), "Modern diagnosis names disease by APPEARANCE",
         "apple scab · rust · Septoria leaf spot · early and late blight · powdery mildew · leaf mould · black rot",
         BLUE)

    _box(c(1, 0), "…and describe symptoms as",
         "paleness, yellowness, drying, withering, shedding, dieback, falling bark, oozing, failure to flower or fruit",
         CREAM); _borders(c(1, 0), on=True, color=brown)
    _box(c(1, 2), "…and describe symptoms as",
         "spots with pale centres, pustules, corky patches, powdery coating, fuzzy mould, sunken decaying areas",
         BLUE)

    overlap = c(0, 1).merge(c(1, 1))
    _box(overlap, "← Overlap →",
         "yellowing · drying · insect damage · poor soil · wounds\n\nanswerable", GREEN)
    _borders(overlap, on=True, color="2F6F4E")

    note = c(2, 0).merge(c(2, 2))
    _blank(note)
    _write(note, "", "No passage in any of the seven sources describes a lesion’s appearance: "
           "11 of 17 disease-name questions have no answer.", bsize=8.5,
           color=RGBColor(0x9B, 0x3B, 0x3B))
    note.paragraphs[0].runs[0].italic = True
    for k in range(3):
        _blank(c(3, k))
    t._tbl.remove(t.rows[3]._tr)          # the 4th row only existed for cell(2,·) merging safety

    _set_widths(t, [6.2, 3.6, 6.2])
    return t


# --------------------------------------------------------------------------- assembly
def main() -> int:
    tmp = HERE / "_chapter_tmp.docx"
    convert(MD, tmp)
    doc = Document(str(tmp))

    pics = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
    builders = [figure1, figure2]
    if len(pics) < len(builders):
        print(f"expected {len(builders)} picture paragraphs, found {len(pics)}")
        return 1
    for p, build in zip(pics, builders):
        table = build(doc)                 # appended at the end of the body...
        p._p.addprevious(table._tbl)       # ...then moved to where the picture was
        p._p.getparent().remove(p._p)      # picture paragraph goes; caption stays
    # The third picture (Figure 3, the results chart) is left in place as a PNG for now;
    # the Word step below replaces it with a native Word chart.

    doc.save(str(OUT)); tmp.unlink(missing_ok=True)
    d = Document(str(OUT))
    print(f"wrote {OUT.relative_to(ROOT)}: {len(d.paragraphs)} paragraphs, "
          f"{len(d.tables)} tables, {len(d.inline_shapes)} images")
    return insert_native_chart()


# --------------------------------------------------------------------------- Figure 3
# Stage data for the results chart. Kept here, next to the builder, so the chart and the
# text cannot drift apart: these are the numbers in Table 2 and EXPERIMENT_LOG.md 6m-6p.
STAGES = [
    ("Stage 1: 233 passages,\nhand-written wording,\nduplicate labels", 13.6, 81.8),
    ("Stage 2: 233 passages,\nsystem's own wording,\nunique labels", 16.7, 83.3),
    ("Stage 3 (final): 270 passages,\n+ Vishvavallabha,\ncitation format fixed", 51.9, 40.7),
]
SERIES = ("answers backed by a real citation", "answerable questions refused")


def _rgb(r, g, b) -> int:
    return r + g * 256 + b * 65536


def insert_native_chart() -> int:
    """Replace the Figure 3 picture with a real Word chart, using Word itself.

    python-docx cannot create charts. Word can, through COM: the chart is a genuine Word
    object with an embedded workbook, so a reviewer who clicks it gets "Edit Data" and
    the numbers, exactly as if it had been made in Word by hand.
    """
    try:
        import win32com.client  # noqa: PLC0415
    except ImportError:
        print("pywin32 not available - Figure 3 left as a picture")
        return 0
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(OUT))
        # python-docx's default template is flagged as an older format, so Word opens it
        # in "Compatibility Mode". Upgrade it: charts behave as in a current document and
        # the reviewer does not see the compatibility banner in the title bar.
        if doc.CompatibilityMode < 15:
            doc.Convert()
        if doc.InlineShapes.Count < 1:
            print("no picture found to replace"); doc.Close(False); return 1
        pic = doc.InlineShapes(1)              # figures 1-2 are tables, so this is fig 3
        rng = pic.Range
        pic.Delete()
        shape = doc.InlineShapes.AddChart2(-1, 51, rng)   # 51 = xlColumnClustered
        chart = shape.Chart

        chart.ChartData.Activate()
        wb = chart.ChartData.Workbook
        ws = wb.Worksheets(1)
        ws.Cells(1, 1).Value = ""
        ws.Cells(1, 2).Value = SERIES[0]
        ws.Cells(1, 3).Value = SERIES[1]
        for i, (label, grounded, refused) in enumerate(STAGES, start=2):
            ws.Cells(i, 1).Value = label
            ws.Cells(i, 2).Value = grounded
            ws.Cells(i, 3).Value = refused
        # A string reference, not the Excel Range object: the Range lives in the embedded
        # workbook's process and Word's Chart.SetSourceData rejects it across that boundary.
        chart.SetSourceData(f"='{ws.Name}'!$A$1:$C${len(STAGES) + 1}")
        wb.Close()

        chart.HasTitle = False
        chart.HasLegend = True
        chart.Legend.Position = -4107          # xlLegendPositionBottom
        chart.Legend.Font.Size = 9
        chart.Legend.Font.Name = FONT
        vax = chart.Axes(2)                    # xlValue
        vax.MinimumScale = 0; vax.MaximumScale = 100; vax.MajorUnit = 20
        vax.HasTitle = True
        vax.AxisTitle.Text = "per cent of answerable questions"
        vax.AxisTitle.Font.Size = 9; vax.AxisTitle.Font.Name = FONT
        vax.TickLabels.Font.Size = 9; vax.TickLabels.Font.Name = FONT
        vax.HasMajorGridlines = True
        vax.MajorGridlines.Format.Line.ForeColor.RGB = _rgb(220, 220, 220)
        cax = chart.Axes(1)                    # xlCategory
        cax.TickLabels.Font.Size = 8.5; cax.TickLabels.Font.Name = FONT
        colours = (_rgb(31, 78, 121), _rgb(201, 162, 39))
        for k in (1, 2):
            s = chart.SeriesCollection(k)
            s.Format.Fill.ForeColor.RGB = colours[k - 1]
            s.HasDataLabels = True
            s.DataLabels().NumberFormat = '0.0"%"'
            s.DataLabels().Font.Size = 8.5
            s.DataLabels().Font.Name = FONT
        chart.ChartGroups(1).GapWidth = 60
        shape.Width = 430; shape.Height = 260  # points
        doc.Save()
        doc.Close(False)
        print("Figure 3 inserted as a native Word chart")
        return 0
    finally:
        word.Quit()


if __name__ == "__main__":
    sys.exit(main())
